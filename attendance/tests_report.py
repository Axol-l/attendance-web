"""
报表生成测试（Phase 4）

只断言"结构 + 口径 + 样式"这些可复现的东西：
  · Sheet 名、标题、表头、人员块与行数
  · 列位置（应出勤/出勤打卡/实际出勤/请假/缺勤/加班）
  · Excel 公式（实际出勤）
  · 状态填充色与周末红字的**色值**（不是数量 —— 数量依赖数据规模）
  · 权限边界
"""
import os
import shutil
import tempfile
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import UserPermission

from . import mapping
from .models import AttendanceDaily, AttendanceRule, LeaveRecord, UploadedFile
from .report import (
    SHEET_DAILY,
    SHEET_SUMMARY,
    ReportError,
    argb,
    build_report,
    default_output_name,
    template_path,
    trim_trailing_columns,
)


class TemplatePresenceTests(TestCase):
    def test_template_is_shipped(self):
        """报表模板必须随代码一起提供 —— 缺了就没法生成"""
        path = template_path()
        self.assertTrue(os.path.exists(path), f'缺少模板 {path}')
        self.assertGreater(os.path.getsize(path), 10_000)

    def test_template_has_expected_sheets(self):
        wb = openpyxl.load_workbook(template_path())
        try:
            self.assertIn(SHEET_DAILY, wb.sheetnames)
            self.assertIn(SHEET_SUMMARY, wb.sheetnames)
            # 表头文案必须对得上，否则说明模板被换过
            ws = wb[SHEET_SUMMARY]
            self.assertEqual(ws.cell(2, 2).value, '姓名')
            self.assertEqual(ws.cell(3, 6).value, '应出勤')
            self.assertEqual(ws.cell(3, 7).value, '出勤打卡')
            self.assertEqual(ws.cell(3, 8).value, '实际出勤')
            self.assertEqual(ws.cell(3, 9).value, '事假')
            self.assertEqual(ws.cell(3, 17).value, '产检假')
            self.assertEqual(ws.cell(3, 18).value, '迟到')
            # 加班在 W2（与 W3 合并），单位 'H' 在 W4
            self.assertEqual(ws.cell(2, 23).value, '加班')
            self.assertEqual(ws.cell(4, 23).value, 'H')
            self.assertIn('W2:W3', {str(x) for x in ws.merged_cells.ranges})
            ws1 = wb[SHEET_DAILY]
            self.assertEqual(ws1.cell(2, 1).value, '姓名')
            self.assertEqual(ws1.cell(2, 14).value, '总计加班/H')
        finally:
            wb.close()

    def test_argb_normalizes_six_digit_colors(self):
        """
        ⚠️ 6 位色值必须补成 8 位 ARGB。
           否则 openpyxl 会把 'FF8080' 当 ARGB 解析、alpha 变 0，
           生成出 '00FF8080'，Excel 里显示异常。
        """
        self.assertEqual(argb('FF8080'), 'FFFF8080')
        self.assertEqual(argb('#ff8080'), 'FFFF8080')
        self.assertEqual(argb('FFFF8080'), 'FFFF8080')
        self.assertIsNone(argb(None))
        self.assertIsNone(argb('xyz'))


class ReportFixtureMixin:
    """造两个月的数据，覆盖正常日、周末、请假、迟到、缺卡、负加班"""

    def setUp(self):
        self.user = User.objects.create_user('reporter', password='x')
        # ⚠️ 报表相关页面要求 attendance.report；
        #    普通用户必须显式授权（管理员才自动拥有全部权限）。
        #    漏了这句会让所有报表请求 302 到首页，而断言 200 的测试就会以
        #    "内容类型不对"这种间接方式失败。
        UserPermission.objects.create(user=self.user,
                                      permission_code='attendance.report')
        self.rule = AttendanceRule.get_active()
        self.tmpdir = tempfile.mkdtemp(prefix='attendance-report-test-')
        self.upload = UploadedFile.objects.create(
            original_filename='daily.xlsx', stored_filename='d.xlsx', file_path='/tmp/d.xlsx',
            file_size=1, file_kind='daily', period='2014-08', uploaded_by=self.user,
        )

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _day(self, day, name='甲', **kw):
        """默认造一条'正常上班'的记录"""
        defaults = dict(
            in1_time='09:15', in1_result='正常', out1_time='18:15', out1_result='正常',
            shift='默认班次 09:15-18:15', work_minutes=Decimal('540'),
            attend_days=Decimal('1'),
        )
        defaults.update(kw)
        return AttendanceDaily.objects.create(
            work_date=date(2014, 8, day), user_id=f'U-{name}', name=name,
            department='测试部', position='工程师',
            source_file=self.upload, created_by=self.user, **defaults,
        )

    def _out(self, name='r.xlsx'):
        return os.path.join(self.tmpdir, name)


class BuildReportStructureTests(ReportFixtureMixin, TestCase):
    def test_creates_file_with_two_sheets(self):
        self._day(1)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        self.assertTrue(os.path.exists(path))
        wb = openpyxl.load_workbook(path)
        try:
            self.assertEqual(wb.sheetnames, [SHEET_DAILY, SHEET_SUMMARY])
        finally:
            wb.close()

    def test_daily_sheet_header_and_first_row(self):
        self._day(1)
        self._day(2, in1_time='09:17')
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            self.assertEqual(ws.cell(1, 1).value, 'jointelli原始考勤数据')
            headers = [ws.cell(2, c).value for c in range(1, 15)]
            self.assertEqual(headers, [
                '姓名', '部门', '职位', '日期', '班次',
                '上班打卡时间', '上班打卡结果', '下班打卡时间', '下班打卡结果',
                '迟到时长', '早退时长', '上班缺卡次数', '下班缺卡次数', '总计加班/H'])
            self.assertEqual(ws.cell(3, 1).value, '甲')
            self.assertEqual(ws.cell(3, 2).value, '测试部')
            self.assertEqual(ws.cell(3, 4).value, '14-08-01 星期五')
            self.assertEqual(ws.cell(4, 4).value, '14-08-02 星期六')
        finally:
            wb.close()

    def test_one_block_per_person_with_vertical_merges(self):
        for d in range(1, 6):
            self._day(d, name='乙')      # 5 天
        for d in range(1, 4):
            self._day(d, name='甲')      # 3 天

        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            blocks = [r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value]
            # 按 Unicode 码点排序：乙(4E59) < 甲(7532)
            # 乙 5 天（3-7），甲 3 天（8-10）
            self.assertEqual(blocks, [3, 8])
            merged = {str(x) for x in ws.merged_cells.ranges}
            # A/B/C 与 J..N 都要纵向合并
            self.assertIn('A3:A7', merged)
            self.assertIn('B3:B7', merged)
            self.assertIn('C3:C7', merged)
            self.assertIn('N3:N7', merged)
            self.assertIn('A8:A10', merged)
            # 最后一行的日期不合并
            self.assertNotIn('D3:D7', merged)
        finally:
            wb.close()

    def test_blocks_do_not_overlap_or_leave_gaps(self):
        """块必须首尾相接，不能重叠也不能有空档"""
        for d in range(1, 4):
            self._day(d, name='乙')      # 3 天
        for d in range(1, 3):
            self._day(d, name='甲')      # 2 天
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            blocks = [r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value]
            self.assertEqual(blocks, [3, 6])       # 乙 3-5，甲 6-7
            merged = {str(x) for x in ws.merged_cells.ranges}
            self.assertIn('A3:A5', merged)
            self.assertIn('A6:A7', merged)
            # 数据只到第 7 行，第 8 行及以后必须为空
            self.assertTrue(all(ws.cell(8, c).value is None for c in range(1, 15)))
        finally:
            wb.close()

    def test_no_extra_rows_beyond_data(self):
        """行数必须正好等于人数×天数 —— 桌面版在这里会多出空行/丢行"""
        for name, days in (('甲', 31), ('乙', 31)):
            for d in range(1, 32):
                self._day(d, name=name)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            data_rows = [r for r in range(3, ws.max_row + 1)
                         if any(ws.cell(r, c).value is not None for c in range(1, 15))]
            self.assertEqual(len(data_rows), 62)
            self.assertEqual(data_rows[-1], 64)
        finally:
            wb.close()

    def test_statistics_columns_are_per_person_totals(self):
        """J..N 是"该人整月合计"，写在块首行，其余由合并覆盖"""
        self._day(1, late_minutes=Decimal('10'), missing_in_count=1)
        self._day(2, late_minutes=Decimal('5'), missing_out_count=2)
        self._day(3, early_leave_minutes=Decimal('20'))

        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            self.assertEqual(ws.cell(3, 10).value, 15.0)    # 迟到合计
            self.assertEqual(ws.cell(3, 11).value, 20.0)    # 早退合计
            self.assertEqual(ws.cell(3, 12).value, 1)       # 上班缺卡
            self.assertEqual(ws.cell(3, 13).value, 2)       # 下班缺卡
        finally:
            wb.close()

    def test_overtime_column_respects_clamping(self):
        """540 分钟 → (540-450)/60 = 1.5 小时；300 分钟 → 截断为 0"""
        self._day(1, work_minutes=Decimal('540'))
        self._day(2, work_minutes=Decimal('300'))
        self._day(3, work_minutes=None)

        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            self.assertEqual(ws.cell(3, 14).value, 1.5)
        finally:
            wb.close()

    def test_overtime_keeps_negative_when_rule_allows(self):
        self.rule.clamp_negative_overtime = False
        self.rule.save()
        self._day(1, work_minutes=Decimal('300'))       # -2.5
        self._day(2, work_minutes=Decimal('540'))       # +1.5
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            self.assertAlmostEqual(wb[SHEET_DAILY].cell(3, 14).value, -1.0, places=6)
        finally:
            wb.close()

    def test_empty_period_raises_readable_error(self):
        with self.assertRaises(ReportError) as ctx:
            build_report('2014-08', rule=self.rule, path=self._out())
        self.assertIn('没有日考勤数据', str(ctx.exception))


class ReportSummarySheetTests(ReportFixtureMixin, TestCase):
    def test_summary_columns_and_formula(self):
        self._day(1, attend_days=Decimal('24'))
        LeaveRecord.objects.create(
            approval_no='L1', applicant_name='甲', leave_type='事假',
            start_time=timezone.make_aware(timezone.datetime(2014, 8, 1, 9, 0)),
            duration_hours=Decimal('15'), duration_days=Decimal('2'),
            is_approved=True, period='2014-08',
            source_file=self.upload, created_by=self.user,
        )
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_SUMMARY]
            self.assertEqual(ws.cell(1, 1).value, 'jointelli-08月考勤表')
            self.assertEqual(ws.cell(5, 1).value, 1)            # 序号
            self.assertEqual(ws.cell(5, 2).value, '甲')
            self.assertEqual(ws.cell(5, 3).value, '测试部')
            # F 应出勤 = 规则值（桌面版误写进了 G 列）
            self.assertEqual(ws.cell(5, 6).value, float(self.rule.monthly_standard_days))
            # G 出勤打卡 = 源文件累计
            self.assertEqual(ws.cell(5, 7).value, 24.0)
            # H 实际出勤 = Excel 公式
            self.assertEqual(ws.cell(5, 8).value, '=G5-I5')
            # I 事假 = 15 小时 ÷ 7.5 = 2 天
            self.assertEqual(ws.cell(5, 9).value, 2.0)
        finally:
            wb.close()

    def test_summary_includes_all_leave_columns(self):
        """9 类请假列都必须能被写入（列位置不能错位）"""
        self._day(1)
        start = timezone.make_aware(timezone.datetime(2014, 8, 1, 9, 0))
        for i, (code, _label) in enumerate(mapping.REPORT_LEAVE_COLUMNS):
            LeaveRecord.objects.create(
                approval_no=f'A{i}', applicant_name='甲', leave_type=code,
                start_time=start, duration_hours=Decimal('7.5'),
                duration_days=Decimal('1'), is_approved=True, period='2014-08',
                source_file=self.upload, created_by=self.user,
            )
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_SUMMARY]
            for i, (code, _label) in enumerate(mapping.REPORT_LEAVE_COLUMNS):
                cell = ws.cell(5, 9 + i)
                self.assertEqual(cell.value, 1.0,
                                 f'{code} 应写入第 {9 + i} 列（{cell.coordinate}）')
            # 请假列共 9 列：I..Q
            self.assertEqual(len(mapping.REPORT_LEAVE_COLUMNS), 9)
        finally:
            wb.close()

    def test_rejected_leave_excluded_by_default(self):
        self._day(1)
        start = timezone.make_aware(timezone.datetime(2014, 8, 1, 9, 0))
        LeaveRecord.objects.create(
            approval_no='R1', applicant_name='甲', leave_type='调休',
            start_time=start, duration_hours=Decimal('7.5'), duration_days=Decimal('1'),
            is_approved=False, period='2014-08',          # 未通过
            source_file=self.upload, created_by=self.user,
        )
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            # 调休在第 10 列，未通过则不应写入
            self.assertIsNone(wb[SHEET_SUMMARY].cell(5, 10).value)
        finally:
            wb.close()

    def test_late_early_absent_and_missing_columns(self):
        self._day(1, late_minutes=Decimal('127'), early_leave_minutes=Decimal('73'),
                  absenteeism_days=Decimal('2'), missing_in_count=4, missing_out_count=7)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_SUMMARY]
            self.assertEqual(ws.cell(5, 18).value, 127.0)   # R 迟到(分)
            self.assertEqual(ws.cell(5, 19).value, 73.0)    # S 早退(分)
            self.assertEqual(ws.cell(5, 20).value, 2.0)     # T 旷工(天)
            self.assertEqual(ws.cell(5, 21).value, 4)       # U 上班忘打卡(次)
            self.assertEqual(ws.cell(5, 22).value, 7)       # V 下班忘打卡(次)
        finally:
            wb.close()

    def test_one_row_per_person_sorted(self):
        for name in ('丙', '甲', '乙'):
            self._day(1, name=name)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_SUMMARY]
            names = [ws.cell(r, 2).value for r in range(5, 8)]
            # 按 Unicode 码点排序：丙(4E19) < 乙(4E59) < 甲(7532)
            self.assertEqual(names, ['丙', '乙', '甲'])
            seqs = [ws.cell(r, 1).value for r in range(5, 8)]
            self.assertEqual(seqs, [1, 2, 3])
        finally:
            wb.close()

    def test_summary_rows_match_daily_people(self):
        """Sheet2 的人数必须与 Sheet1 的人块数一致（桌面版在这里不一致）"""
        for name in ('甲', '乙', '丙'):
            for d in range(1, 4):
                self._day(d, name=name)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws1, ws2 = wb[SHEET_DAILY], wb[SHEET_SUMMARY]
            blocks = {ws1.cell(r, 1).value for r in range(3, ws1.max_row + 1)
                      if ws1.cell(r, 1).value}
            rows = {ws2.cell(r, 2).value for r in range(5, ws2.max_row + 1)
                    if ws2.cell(r, 2).value}
            self.assertEqual(blocks, rows)
            self.assertEqual(len(rows), 3)
        finally:
            wb.close()


class ReportStylingTests(ReportFixtureMixin, TestCase):
    def test_status_fill_colors_match_config(self):
        """状态色必须来自规则配置，且是 8 位 ARGB"""
        self._day(1, in1_result='请假')
        self._day(2, in1_result='出差')
        self._day(3, in1_result='迟到')
        self._day(4, in1_result='缺卡')
        self._day(5, in1_result='外勤')
        self._day(6, in1_result='外出')
        self._day(7, in1_result='补卡审批通过')

        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            expected = {
                3: 'FFFFC000',   # 请假
                4: 'FF00B0F0',   # 出差
                5: 'FFCCFFCC',   # 迟到
                6: 'FFFF8080',   # 缺卡
                7: 'FFFFFF00',   # 外勤
                8: 'FFFFCC99',   # 外出
                9: 'FFFFCC99',   # 补卡审批通过
            }
            for row, want in expected.items():
                got = ws.cell(row, 7).fill.fgColor.rgb      # G 列是上班打卡结果
                self.assertEqual(got, want, f'第 {row} 行 {ws.cell(row, 7).value} 颜色不对')
        finally:
            wb.close()

    def test_weekend_dates_are_red(self):
        """周六/周日的日期标红 —— 用日期星期判断，不靠文本结尾"""
        # 2014-08-01 是周五，02 周六，03 周日，04 周一
        for d in range(1, 5):
            self._day(d)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            def is_red(r):
                c = ws.cell(r, 4)
                return bool(c.font and c.font.color and c.font.color.rgb == 'FFFF0000')
            self.assertFalse(is_red(3), '周五不该标红')
            self.assertTrue(is_red(4), '周六该标红')
            self.assertTrue(is_red(5), '周日该标红')
            self.assertFalse(is_red(6), '周一不该标红')
        finally:
            wb.close()

    def test_fill_survives_vertical_merge(self):
        """
        ⚠️ 回归：必须先合并再上色。
           openpyxl 的 merge_cells 会重置区域内样式，顺序反了填充色会被清掉。
        """
        self._day(1, in1_result='请假', missing_in_count=0)
        self._day(2, in1_result='请假')
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb[SHEET_DAILY]
            # 两天都是"请假"，G 列（结果）按状态上色；J..N 是合并列
            self.assertEqual(ws.cell(3, 7).fill.fgColor.rgb, 'FFFFC000')
            self.assertEqual(ws.cell(4, 7).fill.fgColor.rgb, 'FFFFC000')
            # 合并列 J3:J4 的填充也应存在（模板已预置样式）
            self.assertIn('J3:J4', {str(x) for x in ws.merged_cells.ranges})
        finally:
            wb.close()

    def test_color_map_is_configurable(self):
        """规则里改了颜色，报表要跟着变"""
        self.rule.color_map = {'请假': 'FF00FF00'}      # 改成绿色
        self.rule.save()
        self._day(1, in1_result='请假')
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            self.assertEqual(wb[SHEET_DAILY].cell(3, 7).fill.fgColor.rgb, 'FF00FF00')
        finally:
            wb.close()


class ReportDownloadViewTests(ReportFixtureMixin, TestCase):
    def test_anonymous_redirected(self):
        resp = self.client.get(reverse('attendance:report_generate'))
        self.assertEqual(resp.status_code, 302)

    def test_get_shows_form(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse('attendance:report_generate'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '生成并下载')

    def test_post_generates_and_returns_json(self):
        """
        ⚠️ 生成接口返回 **JSON**，不是文件本体。

        原实现直接回传 `Content-Disposition: attachment` 的响应，
        浏览器开始下载但主页面收不到任何事件 → "正在生成…"永远挂着。
        改成两步后，POST 的 200/JSON 就是可靠的完成信号。
        """
        self._day(1)
        self.client.force_login(self.user)
        resp = self.client.post(reverse('attendance:report_generate'),
                                {'period': '2014-08'})
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data['success'])
        self.assertEqual(data['filename'], '考勤表201408.xlsx')
        self.assertGreater(data['size'], 0)
        self.assertIn('download=1', data['url'])
        self.assertIn('period=2014-08', data['url'])

    def test_download_url_returns_xlsx_attachment(self):
        """生成后再用 download=1 取文件，此时才回传附件"""
        self._day(1)
        self.client.force_login(self.user)
        gen = self.client.post(reverse('attendance:report_generate'),
                               {'period': '2014-08'}).json()
        resp = self.client.get(gen['url'])
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment', resp['Content-Disposition'])
        # 中文文件名按 RFC 5987 百分号编码（filename*=utf-8''...）
        from urllib.parse import unquote
        self.assertIn('考勤表201408.xlsx', unquote(resp['Content-Disposition']))

        import io
        # FileResponse 是流式响应，要用 streaming_content 而不是 content
        blob = b''.join(resp.streaming_content)
        wb = openpyxl.load_workbook(io.BytesIO(blob))
        try:
            self.assertEqual(wb.sheetnames, [SHEET_DAILY, SHEET_SUMMARY])
        finally:
            wb.close()

    def test_download_without_generating_is_rejected(self):
        """没生成过就点下载 → 提示重新生成，而不是 500"""
        self.client.force_login(self.user)
        resp = self.client.get(reverse('attendance:report_generate'),
                              {'period': '2014-08', 'download': '1'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '不存在或已过期')

    def test_regenerating_replaces_previous_file(self):
        """同一会话再次生成时旧文件被清掉，不会在临时目录堆积"""
        self._day(1)
        self.client.force_login(self.user)
        first = self.client.post(reverse('attendance:report_generate'),
                                 {'period': '2014-08'}).json()
        second = self.client.post(reverse('attendance:report_generate'),
                                  {'period': '2014-08'}).json()
        # 两次的下载地址一致（同名文件被覆盖），且都能取到内容
        self.assertEqual(first['url'], second['url'])
        self.assertEqual(self.client.get(second['url']).status_code, 200)

    def test_post_writes_audit_log(self):
        from accounts.models import AuditLog
        self._day(1)
        self.client.force_login(self.user)
        self.client.post(reverse('attendance:report_generate'), {'period': '2014-08'})
        log = AuditLog.objects.filter(action='ATTENDANCE_REPORT_GENERATE').first()
        self.assertIsNotNone(log, '生成报表必须写审计日志')
        self.assertIn('2014-08', log.description)

    def test_missing_period_is_reported(self):
        self.client.force_login(self.user)
        resp = self.client.post(reverse('attendance:report_generate'), {'period': ''})
        self.assertEqual(resp.status_code, 400)
        self.assertFalse(resp.json()['success'])
        self.assertIn('请选择账期', resp.json()['message'])

    def test_empty_period_is_reported(self):
        self.client.force_login(self.user)
        resp = self.client.post(reverse('attendance:report_generate'),
                                {'period': '2014-08'})
        self.assertEqual(resp.status_code, 400)
        self.assertIn('没有日考勤数据', resp.json()['message'])

    def test_requires_report_permission(self):
        viewer = User.objects.create_user('viewer2', password='x')
        UserPermission.objects.create(user=viewer, permission_code='attendance.query')
        self.client.force_login(viewer)
        self.assertEqual(
            self.client.get(reverse('attendance:report_generate')).status_code, 302)
        self.assertEqual(
            self.client.post(reverse('attendance:report_generate'),
                             {'period': '2014-08'}).status_code, 302)

    def test_default_output_name_uses_period(self):
        """文件名带账期，避免不同月份互相覆盖"""
        self.assertEqual(default_output_name('2014-08'), '考勤表201408.xlsx')
        self.assertEqual(default_output_name('2014-12'), '考勤表201412.xlsx')

    def test_get_without_period_shows_form(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse('attendance:report_generate'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '生成并下载')

    def test_report_page_script_has_no_serverside_leftovers(self):
        """
        ⚠️ 回归：内联 script 里不能出现未渲染的模板标记，
           也不能把开发期注释（`//` 说明）发给浏览器 ——
           用户反馈"页面出现了无关的注释"。
        """
        self.client.force_login(self.user)
        html = self.client.get(reverse('attendance:report_generate')).content.decode()
        for marker in ('{%', '{{', '{#'):
            self.assertNotIn(marker, html, f'页面里残留了未渲染的 {marker}')

        import re
        scripts = re.findall(r'<script>(.*?)</script>', html, re.S)
        body = '\n'.join(s for s in scripts if 'reportForm' in s)
        self.assertTrue(body, '找不到报表页的脚本')
        # 允许分隔线注释，但不允许成段的中文说明
        prose = [ln.strip() for ln in body.splitlines()
                 if ln.strip().startswith('//') and len(ln.strip()) > 30]
        self.assertEqual(prose, [], f'脚本里还有开发期注释：{prose[:3]}')


class ColumnTrimmingTests(ReportFixtureMixin, TestCase):
    """
    ⚠️ 回归：模板的 `<cols>` 里有铺到最大列的默认列宽定义
       （sheet1 `max=16384`、sheet2 `max=16377`），
       Excel 会据此渲染出整整 16384 列（用户看到的就是 XFD 列）。
       生成时必须把这些尾部定义截断/删除。
    """

    def _cols_ranges(self, path, sheet_index):
        import re
        import zipfile
        with zipfile.ZipFile(path) as z:
            xml = z.read(f'xl/worksheets/sheet{sheet_index}.xml').decode('utf-8')
        block = re.search(r'<cols>(.*?)</cols>', xml, re.S)
        if not block:
            return []
        return [(int(a), int(b)) for a, b in
                re.findall(r'min="(\d+)"\s+max="(\d+)"', block.group(1))]

    def test_no_column_definition_extends_to_max_column(self):
        self._day(1)
        path = build_report('2014-08', rule=self.rule, path=self._out())

        for idx, limit, label in ((1, 15, 'Sheet1'), (2, 25, 'Sheet2')):
            ranges = self._cols_ranges(path, idx)
            worst = max((hi for _lo, hi in ranges), default=0)
            self.assertLessEqual(worst, limit,
                                 f'{label} 的列宽定义延伸到第 {worst} 列（应 ≤ {limit}）')

    def test_column_widths_are_preserved(self):
        """截断只能去掉"数据列之外"的定义，不能丢掉真实列宽"""
        self._day(1)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws1, ws2 = wb[SHEET_DAILY], wb[SHEET_SUMMARY]
            # 模板里 A=15.45 / C=24 / E=28.55（Sheet1），A=5.5 / X=35.18（Sheet2）
            self.assertAlmostEqual(ws1.column_dimensions['A'].width, 15.45, places=1)
            self.assertAlmostEqual(ws1.column_dimensions['C'].width, 24.0, places=1)
            self.assertAlmostEqual(ws2.column_dimensions['A'].width, 5.5, places=1)
            self.assertAlmostEqual(ws2.column_dimensions['X'].width, 35.18, places=1)
        finally:
            wb.close()

    def test_dimension_and_max_column_are_sane(self):
        self._day(1)
        path = build_report('2014-08', rule=self.rule, path=self._out())
        wb = openpyxl.load_workbook(path)
        try:
            ws1, ws2 = wb[SHEET_DAILY], wb[SHEET_SUMMARY]
            self.assertEqual(ws1.max_column, 15)
            self.assertEqual(ws2.max_column, 25)
            self.assertEqual(ws1.calculate_dimension().split(':')[1][0], 'O')
        finally:
            wb.close()

    def test_shipped_template_itself_has_the_problem(self):
        """
        记录事实：**随包的模板本身就带这个缺陷**。
        本测试不是要求模板干净（它是客户提供的资产），
        而是防止有人"顺手"把模板换成一个同样铺满列的版本却不加处理。
        """
        ranges = self._cols_ranges(template_path(), 1)
        worst = max((hi for _lo, hi in ranges), default=0)
        self.assertGreater(
            worst, 15,
            '模板里的列宽定义未铺到数据列之外 —— 若模板已换，'
            '请确认 trim_trailing_columns 是否还需要')
